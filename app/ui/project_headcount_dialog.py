from datetime import date

from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont, QFontMetrics
from PySide6.QtWidgets import (
    QAbstractItemView,
    QButtonGroup,
    QComboBox,
    QDialog,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)
from openpyxl import Workbook
from openpyxl.styles import Font as XlFont

from app.config import AppConfig
from app.queries.project_headcount import (
    get_diary_date_bounds,
    get_person_project_month_hours,
    get_project_month_hours,
    period_months,
)
from app.ui.table_utils import NumericTableWidgetItem, enable_header_sorting
from app.ui.theme import POPUP_GRID_FONT_PX, POPUP_HEIGHT, POPUP_WIDTH, current_theme

_ROW_HEIGHT = 28

PROJECT_FIXED_COLUMNS = ["순번", "프로젝트", "투입인원", "합계(H)"]
PERSON_FIXED_COLUMNS = ["순번", "소속", "이름", "프로젝트", "합계(H)"]


def _fmt_hours(value: float) -> str:
    if not value:
        return ""
    if abs(value - round(value)) < 0.05:
        return f"{int(round(value)):,}"
    return f"{value:,.1f}"


class ProjectHeadcountDialog(QDialog):
    def __init__(self, app_config: AppConfig, parent=None):
        super().__init__(parent)
        self._config = app_config
        self._period_months_list: list[tuple[int, int]] = []
        self._project_rows = []
        self._person_rows = []

        self.setWindowTitle("프로젝트 인력 투입")
        self.setWindowFlags(self.windowFlags() | Qt.WindowType.WindowMaximizeButtonHint)
        self.resize(POPUP_WIDTH, POPUP_HEIGHT)

        today = date.today()
        try:
            min_dt, max_dt = get_diary_date_bounds(self._config.database)
        except Exception:  # noqa: BLE001 - 조회 실패 시 현재연도 기준으로 폴백
            min_dt, max_dt = None, None
        anchor = max_dt or today
        self._end_year, self._end_month = anchor.year, anchor.month
        start_anchor_y = anchor.year - 1
        self._start_year, self._start_month = start_anchor_y, anchor.month % 12 + 1
        # 연도 콤보 범위(min_year~max_year)에 위에서 계산한 시작연도(start_anchor_y)가
        # 반드시 포함돼야 한다 — 안 그러면 아래 _year_combo의 findData가 실패해서
        # "시작"이 콤보 마지막 값(최댓값 연도)으로 튀어버려, 조회기간이 "미래 ~ 과거"
        # 처럼 거꾸로 표시되는 버그가 있었다.
        self._min_year = min(start_anchor_y, min_dt.year if min_dt else today.year)
        self._max_year = max(anchor.year, today.year) + 1

        title_label = QLabel("프로젝트 인력 투입")
        title_label.setProperty("role", "title")

        self._project_radio = QRadioButton("프로젝트 기준")
        self._project_radio.setChecked(True)
        self._person_radio = QRadioButton("사람 기준")
        view_group = QButtonGroup(self)
        view_group.addButton(self._project_radio)
        view_group.addButton(self._person_radio)
        self._project_radio.toggled.connect(self._on_view_changed)

        def _year_combo(selected: int, fallback_to_latest: bool) -> QComboBox:
            """selected 연도가 범위 밖이라 못 찾으면(원래는 안 생겨야 하지만 방어적으로),
            시작 콤보는 최솟값, 종료 콤보는 최댓값으로 폴백한다 — 둘 다 최댓값으로
            폴백하면 "시작이 종료보다 미래"인 거꾸로 된 기간이 나올 수 있다."""
            combo = QComboBox()
            for year in range(self._min_year, self._max_year + 1):
                combo.addItem(f"{year}년", year)
            idx = combo.findData(selected)
            if idx < 0:
                idx = combo.count() - 1 if fallback_to_latest else 0
            combo.setCurrentIndex(idx)
            return combo

        def _month_combo(selected: int) -> QComboBox:
            combo = QComboBox()
            for month in range(1, 13):
                combo.addItem(f"{month}월", month)
            combo.setCurrentIndex(selected - 1)
            return combo

        self._start_year_combo = _year_combo(self._start_year, fallback_to_latest=False)
        self._start_month_combo = _month_combo(self._start_month)
        self._end_year_combo = _year_combo(self._end_year, fallback_to_latest=True)
        self._end_month_combo = _month_combo(self._end_month)

        search_button = QPushButton("조회")
        search_button.setProperty("variant", "secondary")
        search_button.clicked.connect(self._on_search_period)

        refresh_button = QPushButton("새로고침")
        refresh_button.setProperty("variant", "secondary")
        refresh_button.clicked.connect(self._load_data)

        export_button = QPushButton("엑셀로 저장")
        export_button.setProperty("variant", "secondary")
        export_button.clicked.connect(self._export_to_excel)

        header_row = QHBoxLayout()
        header_row.addWidget(title_label)
        header_row.addSpacing(20)
        header_row.addWidget(self._project_radio)
        header_row.addWidget(self._person_radio)
        header_row.addSpacing(12)
        header_row.addWidget(QLabel("조회기간:"))
        header_row.addWidget(self._start_year_combo)
        header_row.addWidget(self._start_month_combo)
        header_row.addWidget(QLabel("~"))
        header_row.addWidget(self._end_year_combo)
        header_row.addWidget(self._end_month_combo)
        header_row.addWidget(search_button)
        header_row.addStretch()
        header_row.addWidget(refresh_button)
        header_row.addWidget(export_button)

        self._status_label = QLabel("")
        self._status_label.setProperty("role", "secondary")

        self._project_table = self._make_grid()
        # 투입인원이 2명 이상인 프로젝트는 그 밑에 참여자별 세부 행을 붙이는 구조라
        # (아래 _build_project_table 참고), 헤더 클릭 정렬을 켜면 그 묶음이 깨진다.
        self._person_table = self._make_grid()
        enable_header_sorting(self._person_table)

        self._view_stack = QStackedWidget()
        self._view_stack.addWidget(self._project_table)
        self._view_stack.addWidget(self._person_table)

        layout = QVBoxLayout()
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(10)
        layout.addLayout(header_row)
        layout.addWidget(self._status_label, 0)
        layout.addWidget(self._view_stack, 1)
        self.setLayout(layout)

        self._load_data()

    @staticmethod
    def _make_grid() -> QTableWidget:
        table = QTableWidget()
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        table.setSelectionMode(QAbstractItemView.SelectionMode.NoSelection)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        table.setStyleSheet(f"font-size: {POPUP_GRID_FONT_PX}px;")
        return table

    def _on_view_changed(self, _checked: bool) -> None:
        self._view_stack.setCurrentIndex(0 if self._project_radio.isChecked() else 1)

    def _on_search_period(self) -> None:
        sy = self._start_year_combo.currentData()
        sm = self._start_month_combo.currentData()
        ey = self._end_year_combo.currentData()
        em = self._end_month_combo.currentData()
        if (sy, sm) > (ey, em):
            QMessageBox.warning(self, "조회기간", "시작 연월이 종료 연월보다 늦을 수 없습니다.")
            return
        self._start_year, self._start_month = sy, sm
        self._end_year, self._end_month = ey, em
        self._load_data()

    def _load_data(self) -> None:
        self._period_months_list = period_months(
            self._start_year, self._start_month, self._end_year, self._end_month
        )
        if len(self._period_months_list) > 60:
            QMessageBox.warning(
                self, "조회기간", "조회기간이 너무 깁니다(최대 60개월). 기간을 줄여 주세요."
            )
            return
        try:
            self._project_rows = get_project_month_hours(
                self._config.database, self._start_year, self._start_month, self._end_year, self._end_month
            )
            self._person_rows = get_person_project_month_hours(
                self._config.database, self._start_year, self._start_month, self._end_year, self._end_month
            )
        except Exception as exc:  # noqa: BLE001 - 조회 실패를 사용자에게 그대로 안내
            QMessageBox.critical(self, "조회 실패", f"인력 투입 정보 조회에 실패했습니다.\n{exc}")
            return

        self._build_project_table()
        self._build_person_table()
        self._status_label.setText(
            f"프로젝트 {len(self._project_rows)}건 · (직원×프로젝트) {len(self._person_rows)}행 · "
            f"{self._start_year}.{self._start_month:02d} ~ {self._end_year}.{self._end_month:02d} 기준 "
            f"(단위: 시간, 업무일지 tb_wrkst_diary_info)"
        )

    # ------------------------------------------------------------------
    def _month_headers(self) -> list[str]:
        return [f"{y}.{m}월" for y, m in self._period_months_list]

    def _month_col_width(self, table: QTableWidget) -> int:
        metrics = QFontMetrics(table.horizontalHeader().font())
        widest = max((metrics.horizontalAdvance(h) for h in self._month_headers()), default=0)
        return max(58, min(widest + 22, 110))

    def _set_hours_cell(self, table: QTableWidget, row: int, col: int, value: float, bold: bool = False) -> None:
        item = NumericTableWidgetItem(_fmt_hours(value))
        item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        theme = current_theme()
        if not value:
            item.setForeground(QColor(theme.text_secondary))
        if bold:
            font = QFont()
            font.setBold(True)
            item.setFont(font)
        table.setItem(row, col, item)

    def _set_text_cell(
        self, table: QTableWidget, row: int, col: int, text: str, center: bool = False, numeric: bool = False
    ) -> None:
        item = NumericTableWidgetItem(text) if numeric else QTableWidgetItem(text)
        if center:
            item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        table.setItem(row, col, item)

    # ------------------------------------------------------------------
    def _persons_by_project(self) -> dict[str, list]:
        grouped: dict[str, list] = {}
        for prow in self._person_rows:
            grouped.setdefault(prow.prj_id, []).append(prow)
        return grouped

    def _build_project_table(self) -> None:
        """투입인원이 2명 이상인 프로젝트는 합계 행 바로 아래에 참여자별 세부 행을
        덧붙인다(이름·소속·개인별 월별 투입시간). 1명뿐인 프로젝트는 굳이 똑같은
        내용을 두 번 보여줄 필요가 없어 세부 행을 만들지 않는다."""
        months = self._period_months_list
        fixed = len(PROJECT_FIXED_COLUMNS)
        table = self._project_table
        table.setSortingEnabled(False)
        table.clear()
        table.setColumnCount(fixed + len(months))
        table.setHorizontalHeaderLabels(PROJECT_FIXED_COLUMNS + self._month_headers())
        for col_idx, width in enumerate((50, 320, 80, 90)):
            table.setColumnWidth(col_idx, width)
        month_width = self._month_col_width(table)
        for col in range(fixed, fixed + len(months)):
            table.setColumnWidth(col, month_width)

        persons_by_prj = self._persons_by_project()
        display_rows: list[tuple[str, object]] = []
        for prow in self._project_rows:
            display_rows.append(("project", prow))
            if prow.headcount > 1:
                for person in persons_by_prj.get(prow.prj_id, []):
                    display_rows.append(("person", person))

        theme = current_theme()
        table.setRowCount(len(display_rows))
        seq = 0
        for row_idx, (kind, item) in enumerate(display_rows):
            table.setRowHeight(row_idx, _ROW_HEIGHT)
            if kind == "project":
                seq += 1
                self._set_text_cell(table, row_idx, 0, str(seq), center=True, numeric=True)
                self._set_text_cell(table, row_idx, 1, item.prj_name)
                self._set_text_cell(table, row_idx, 2, str(item.headcount), center=True, numeric=True)
                self._set_hours_cell(table, row_idx, 3, item.total_hours, bold=True)
                for offset, (y, m) in enumerate(months):
                    self._set_hours_cell(
                        table, row_idx, 4 + offset, item.monthly.get(f"{y:04d}{m:02d}", 0.0)
                    )
                continue

            # kind == "person": 참여자 세부 행 — 순번/투입인원 칸은 비우고, 이름
            # 앞에 들여쓰기 표시를 붙여 위 프로젝트 행에 속한다는 걸 나타낸다.
            self._set_text_cell(table, row_idx, 0, "")
            name_item = QTableWidgetItem(f"　└ {item.empl_name} ({item.dept})")
            name_item.setForeground(QColor(theme.text_secondary))
            table.setItem(row_idx, 1, name_item)
            self._set_text_cell(table, row_idx, 2, "")
            hours_item = NumericTableWidgetItem(_fmt_hours(item.total_hours))
            hours_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            hours_item.setForeground(QColor(theme.text_secondary))
            table.setItem(row_idx, 3, hours_item)
            for offset, (y, m) in enumerate(months):
                value = item.monthly.get(f"{y:04d}{m:02d}", 0.0)
                cell_item = NumericTableWidgetItem(_fmt_hours(value))
                cell_item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                if not value:
                    cell_item.setForeground(QColor(theme.text_secondary))
                table.setItem(row_idx, 4 + offset, cell_item)
        # 세부 행이 있는 프로젝트는 그 묶음이 정렬로 흐트러지면 안 되므로, 이 표는
        # 정렬을 켜지 않는다(person_table에서 개인 기준으로 얼마든지 정렬 가능).

    def _build_person_table(self) -> None:
        months = self._period_months_list
        fixed = len(PERSON_FIXED_COLUMNS)
        table = self._person_table
        table.setSortingEnabled(False)
        table.clear()
        table.setColumnCount(fixed + len(months))
        table.setHorizontalHeaderLabels(PERSON_FIXED_COLUMNS + self._month_headers())
        for col_idx, width in enumerate((50, 130, 80, 300, 90)):
            table.setColumnWidth(col_idx, width)
        month_width = self._month_col_width(table)
        for col in range(fixed, fixed + len(months)):
            table.setColumnWidth(col, month_width)

        # 예전엔 같은 직원의 두 번째 프로젝트부터 소속/이름을 비워 시각적으로
        # 묶어줬는데, 헤더 클릭 정렬을 켜면 행 순서가 바뀌면서 그 "묶음"이
        # 아무 의미 없어지므로(엉뚱한 행이 이름 없이 남는다) 매 행에 소속/이름을
        # 그대로 반복 표시하는 쪽으로 바꿨다 — 정렬해도 각 행이 누구 건지 항상
        # 보인다.
        table.setRowCount(len(self._person_rows))
        for row_idx, prow in enumerate(self._person_rows):
            table.setRowHeight(row_idx, _ROW_HEIGHT)
            self._set_text_cell(table, row_idx, 0, str(row_idx + 1), center=True, numeric=True)
            self._set_text_cell(table, row_idx, 1, prow.dept)
            self._set_text_cell(table, row_idx, 2, prow.empl_name)
            self._set_text_cell(table, row_idx, 3, prow.prj_name)
            self._set_hours_cell(table, row_idx, 4, prow.total_hours, bold=True)
            for offset, (y, m) in enumerate(months):
                self._set_hours_cell(
                    table, row_idx, 5 + offset, prow.monthly.get(f"{y:04d}{m:02d}", 0.0)
                )
        table.setSortingEnabled(True)

    # ------------------------------------------------------------------
    def _export_to_excel(self) -> None:
        is_project = self._project_radio.isChecked()
        rows = self._project_rows if is_project else self._person_rows
        if not rows:
            QMessageBox.information(self, "엑셀로 저장", "저장할 데이터가 없습니다.")
            return

        default_name = (
            "프로젝트별_인력투입.xlsx" if is_project else "사람별_프로젝트_인력투입.xlsx"
        )
        path, _ = QFileDialog.getSaveFileName(
            self, "엑셀로 저장", default_name, "Excel Files (*.xlsx)"
        )
        if not path:
            return

        months = self._period_months_list
        month_headers = self._month_headers()
        workbook = Workbook()
        sheet = workbook.active
        zero_font = XlFont(color="9AA0A6")

        if is_project:
            sheet.title = "프로젝트별 인력투입"
            headers = list(PROJECT_FIXED_COLUMNS) + month_headers
            for col_idx, text in enumerate(headers, start=1):
                sheet.cell(row=1, column=col_idx, value=text)

            persons_by_prj = self._persons_by_project()
            person_font = XlFont(color="5B5F73")
            excel_row = 2
            seq = 0
            for prow in self._project_rows:
                seq += 1
                values = [seq, prow.prj_name, prow.headcount, prow.total_hours] + [
                    prow.monthly.get(f"{y:04d}{m:02d}", 0.0) for y, m in months
                ]
                for col_idx, value in enumerate(values, start=1):
                    cell = sheet.cell(row=excel_row, column=col_idx, value=value)
                    if col_idx >= 4 and not value:
                        cell.font = zero_font
                excel_row += 1

                if prow.headcount <= 1:
                    continue
                for person in persons_by_prj.get(prow.prj_id, []):
                    person_values = [None, f"　└ {person.empl_name} ({person.dept})", None, person.total_hours] + [
                        person.monthly.get(f"{y:04d}{m:02d}", 0.0) for y, m in months
                    ]
                    for col_idx, value in enumerate(person_values, start=1):
                        cell = sheet.cell(row=excel_row, column=col_idx, value=value)
                        cell.font = person_font
                    excel_row += 1
        else:
            sheet.title = "사람별 프로젝트 인력투입"
            headers = list(PERSON_FIXED_COLUMNS) + month_headers
            for col_idx, text in enumerate(headers, start=1):
                sheet.cell(row=1, column=col_idx, value=text)
            for row_idx, prow in enumerate(self._person_rows, start=2):
                values = [
                    row_idx - 1,
                    prow.dept,
                    prow.empl_name,
                    prow.prj_name,
                    prow.total_hours,
                ] + [prow.monthly.get(f"{y:04d}{m:02d}", 0.0) for y, m in months]
                for col_idx, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                    if col_idx >= 5 and not value:
                        cell.font = zero_font

        try:
            workbook.save(path)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.critical(self, "저장 실패", f"엑셀 파일 저장에 실패했습니다.\n{exc}")
            return
        QMessageBox.information(self, "엑셀로 저장", f"저장이 완료되었습니다.\n{path}")
